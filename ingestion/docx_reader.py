from docx import Document


def load_docx(path: str) -> str:
    doc = Document(path)
    text = "\n".join(para.text for para in doc.paragraphs)
    return text


def docx_reader(path: str ) -> str:
    doc = Document(path)
    pieces = []
    for paragraph in doc.paragraphs:
        pieces.append(paragraph.text)
    
    text = "\n".join(pieces)

    return text